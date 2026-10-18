#!/usr/bin/env python3
"""
19_ni_score.py - Ipotesi H4: punteggio di invasione neurale (NI score) contro errore di etichetta nel tumore.

Scopo
-----
La Table S1A del supplementare di Chen et al., Cancer Cell 2025 (`mmc2.xlsx`) riporta
i metadati clinici dei pazienti, compreso il punteggio di invasione neurale (`NI
score`). I 14 pazienti con `scRNA-seq = Yes` corrispondono ai 14 pazienti del dataset
single-cell. Ipotesi H4, formulata prima di calcolare la correlazione (stesso testo
stampato dallo script, `H4_TEXT`): NI score piu' alto -> microambiente piu'
immunosoppresso -> piu' T esauste -> piu' errore di annotazione delle T nel tumore
(attesa: rho > 0). H4 passa per "piu' T esauste", quindi eredita i problemi di H1
(16_pseudobulk_patient.py): una deriva da CD8 a CD4T non e' plausibile, la direzione
ipotizzata e' opposta alla maggior parte degli errori osservati (CD4 -> CD8T) e i geni
di esaurimento sono alti anche nelle Treg e nelle CD4 Tfh-like. Il gruppo dicotomico (`NI group`: 11 High contro 3 Low) e' troppo sbilanciato:
si usa il punteggio continuo.

Test: Spearman con permutazione Monte Carlo (20.000 permutazioni casuali, non
esatto) e IC bootstrap sui pazienti ("expanded percentile"). H4 usa lo stesso
outcome di H1-H3 (16_pseudobulk_patient.py): i p grezzi di H1-H3 vengono ricaricati
e Holm-Bonferroni e' ricalcolato sulla famiglia completa di 4 ipotesi.

Input
-----
  --table-s1        data/raw/paper_supplement/mmc2.xlsx (Table S1 di Chen et al. 2025)
  --arbitration     results/tables/arbitration_blood.csv (da 10_loco.py) (pareggi esclusi, come in 16)
  --previous-tests  results/tables/pseudobulk_hypothesis_tests.csv (da 16_pseudobulk_patient.py)

Output
------
  results/tables/ni_score_joined.csv       outcome e variabili cliniche per paziente
  results/tables/ni_score_test.csv         rho, p grezzo e IC di H4
  results/tables/ni_score_holm_joint.csv   p grezzo e p di Holm per H1-H4 (colonna
                                           p_holm_family_of_{n}, n = numero di test con p
                                           definito; 4 con --previous-tests, 1 senza)
  results/figures/ni_score_vs_error.png
  (i pazienti con NI score mancante o non numerico sono esclusi dal test, con i loro
  ID stampati; ni_score_joined.csv li contiene tutti. Con meno di 6 pazienti con NI
  score definito lo script si ferma e cancella test, Holm e figura di un'esecuzione
  precedente)

Argomenti
---------
  --table-s1 (default data/raw/paper_supplement/mmc2.xlsx), --arbitration, --previous-tests

Uso
---
    python src/19_ni_score.py --table-s1 data/raw/paper_supplement/mmc2.xlsx \\
        --arbitration results/tables/arbitration_blood.csv \\
        --previous-tests results/tables/pseudobulk_hypothesis_tests.csv
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import matplotlib

# Configurazione del backend grafico non interattivo per esecuzione batch
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from scpipe import holm_bonferroni, percentile_ci, permutation_corr_test  # noqa: E402

# Percorsi di output e costanti del protocollo inferenziale
FIG = Path("results/figures")
TAB = Path("results/tables")
SEED = 0
N_PERM = 20000
MIN_CLONES_OUTCOME = 5
# Unica formulazione di H4 (docstring e messaggi)
H4_TEXT = (
    "NI score piu' alto -> microambiente piu' immunosoppresso -> piu' T esauste -> "
    "piu' errore di annotazione delle T nel tumore (attesa: rho > 0)."
)

# Schema atteso delle colonne nella Tabella Supplementare S1A (mmc2.xlsx)
EXPECTED_COLS = [
    "PatientID",
    "Age",
    "Gender",
    "Clinical Stage",
    "pTNM",
    "scRNA-seq",
    "NI score",
    "NI group",
]


# --------------------------------------------------------------------------- #
def clean_patient_id(raw: str) -> str:
    """Normalizza l'identificativo grezzo del paziente nel formato canonico PAxx
    a due cifre per garantire l'allineamento con le chiavi di scRNA-seq."""
    m = re.search(r"(\d+)", str(raw))
    if not m:
        raise ValueError(f"ID paziente non riconosciuto: {raw!r}")
    return f"PA{int(m.group(1)):02d}"


def load_table_s1a(path: Path) -> pd.DataFrame:
    """Carica e valida il foglio 'Table S1A' dal file Excel supplementare,
    verificando l'integrita' delle intestazioni e normalizzando gli ID paziente."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Table S1A" not in wb.sheetnames:
        raise SystemExit(
            f"[ERRORE] foglio 'Table S1A' assente. Fogli trovati: {wb.sheetnames}"
        )
    ws = wb["Table S1A"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        raise SystemExit(
            f"[ERRORE] 'Table S1A' ha {len(rows)} righe: servono titolo, intestazione e dati."
        )
    # Estrazione delle intestazioni dalla seconda riga e verifica delle colonne obbligatorie
    header = [str(h).strip() if h else h for h in rows[1]]
    missing = [c for c in EXPECTED_COLS if c not in header]
    if missing:
        raise SystemExit(
            f"[ERRORE] colonne attese assenti: {missing}\n"
            f"colonne trovate: {header}\n"
            "Il formato del supplementare potrebbe essere cambiato: controlla a mano."
        )
    data = [dict(zip(header, r)) for r in rows[2:] if r[0] is not None]
    if not data:
        raise SystemExit(
            "[ERRORE] 'Table S1A' senza righe di dati (solo intestazione)."
        )
    df = pd.DataFrame(data)
    df["patient"] = df.PatientID.apply(clean_patient_id)
    return df


def bootstrap_ci_corr(
    x: np.ndarray, y: np.ndarray, n_boot: int = 2000, seed: int = SEED
):
    """IC 95% di Spearman con bootstrap sui pazienti (righe), "expanded percentile"."""
    rng = np.random.default_rng(seed)
    n = len(x)
    vals = []
    for _ in range(n_boot):
        idx = rng.integers(0, n, n)
        if len(np.unique(x[idx])) < 2 or len(np.unique(y[idx])) < 2:
            continue
        r = np.corrcoef(pd.Series(x[idx]).rank(), pd.Series(y[idx]).rank())[0, 1]
        if np.isfinite(r):
            vals.append(r)
    if len(vals) < 20:
        return np.nan, np.nan
    return percentile_ci(np.array(vals), n)


def build_outcome(arb_path: Path) -> pd.DataFrame:
    """Calcola il tasso medio di errore di annotazione per paziente nel tessuto
    tumorale utilizzando esattamente la medesima soglia di 16_pseudobulk_patient.py."""
    if not arb_path.exists():
        raise SystemExit(f"[ERRORE] {arb_path} assente. Esegui prima 10_loco.py.")
    j = pd.read_csv(arb_path)
    j = j[j.tissue == "Tumor"]
    # unita' con etichetta di maggioranza in pareggio: errore non definito, escluse (come 16)
    n_tie = int(j.error.isna().sum())
    j = j.dropna(subset=["error"])
    print(f"[i] {n_tie} coppie clone-compartimento in pareggio escluse dall'esito")
    out = j.groupby("patient").agg(
        error_rate=("error", "mean"), n_clones=("error", "size")
    )
    out = out[out.n_clones >= MIN_CLONES_OUTCOME]
    return out.reset_index()


# --------------------------------------------------------------------------- #
def main() -> None:
    """Orchestra l'allineamento dei metadati clinici NI score, il test di
    permutazione per H4 e il ricalcolo congiunto di Holm-Bonferroni su H1-H4."""
    ap = argparse.ArgumentParser(
        description="Ipotesi H4: NI score contro errore di etichetta nel tumore.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    ap.add_argument(
        "--table-s1",
        type=Path,
        default=Path("data/raw/paper_supplement/mmc2.xlsx"),
        help="Table S1 del supplemento di Chen et al. 2025 (foglio Table S1A)",
    )
    ap.add_argument(
        "--arbitration",
        type=Path,
        default=Path("results/tables/arbitration_blood.csv"),
        help="identita' clonale e errore per (clone, compartimento), schema 'blood' (da 10_loco.py)",
    )
    ap.add_argument(
        "--previous-tests",
        type=Path,
        default=Path("results/tables/pseudobulk_hypothesis_tests.csv"),
        help="test H1-H3 (da 16_pseudobulk_patient.py) per la correzione di Holm congiunta",
    )
    args = ap.parse_args()

    FIG.mkdir(parents=True, exist_ok=True)
    TAB.mkdir(parents=True, exist_ok=True)

    # Caricamento dei metadati clinici e selezione della coorte sottoposta a scRNA-seq
    if not args.table_s1.exists():
        raise SystemExit(
            f"[ERRORE] {args.table_s1} assente: scarica la Table S1 (mmc2.xlsx) dal "
            "supplementare di Chen et al. 2025 (vedi README)."
        )
    clin = load_table_s1a(args.table_s1)
    scrna = clin[clin["scRNA-seq"] == "Yes"].copy()
    print(f"[ok] pazienti con scRNA-seq=Yes in Table S1A: {len(scrna)}")
    print(
        scrna[["patient", "Age", "Gender", "Clinical Stage", "NI score", "NI group"]]
        .sort_values("patient")
        .to_string(index=False)
    )

    print(
        f"\n[i] NI group: {scrna['NI group'].value_counts().to_dict()} "
        "-> troppo sbilanciato per un test binario, uso solo NI score (continuo)."
    )

    # Join relazionale tra tasso di errore nel tumore e variabili cliniche per paziente
    outcome = build_outcome(args.arbitration)
    df = outcome.merge(
        scrna[["patient", "NI score", "Clinical Stage", "Age"]],
        on="patient",
        how="inner",
    )
    print(f"\n[ok] join outcome x NI score: {len(df)} pazienti -> {sorted(df.patient)}")
    df.to_csv(TAB / "ni_score_joined.csv", index=False)

    print("\n" + "#" * 78)
    print("# IPOTESI H4 (formulata prima di calcolare la correlazione):")
    print(f"# {H4_TEXT}")
    print("#" * 78)

    # Pazienti senza NI score numerico e finito: esclusi dal test, con i loro ID
    ni = pd.to_numeric(df["NI score"], errors="coerce")
    bad = ~np.isfinite(ni.to_numpy(dtype=float))
    if bad.any():
        print(
            f"[!!] NI score mancante o non numerico per {int(bad.sum())} pazienti, "
            f"esclusi dal test: {sorted(df.patient[bad])}"
        )
    df = df.loc[~bad].assign(**{"NI score": ni[~bad]})

    # Test di correlazione per H4: permutazione Monte Carlo e intervallo bootstrap
    if len(df) < 6:
        # niente risultati di un'esecuzione precedente in results/
        for path in (
            TAB / "ni_score_test.csv",
            TAB / "ni_score_holm_joint.csv",
            FIG / "ni_score_vs_error.png",
        ):
            path.unlink(missing_ok=True)
        raise SystemExit(
            f"[ERRORE] solo {len(df)} pazienti con outcome e NI score definito (minimo 6)."
        )
    x, y = df["NI score"].values.astype(float), df.error_rate.values
    rho, p = permutation_corr_test(x, y, N_PERM, SEED)
    lo, hi = bootstrap_ci_corr(x, y)
    print(
        f"\nn={len(df)} pazienti | rho={rho:+.3f}  p_perm={p:.4f}  IC95% boot=[{lo:+.3f}, {hi:+.3f}]"
    )
    if np.isnan(rho):
        print("rho non definito (covariata o esito costante)")
    else:
        print(f"{'COERENTE con H4 (rho>0)' if rho > 0 else 'direzione OPPOSTA a H4'}")
    pd.DataFrame(
        [
            {
                "hypothesis": "H4_ni_score",
                "n_patients": len(df),
                "rho": rho,
                "p_perm": p,
                "ci_lo": lo,
                "ci_hi": hi,
                "n_perm": N_PERM,
            }
        ]
    ).to_csv(TAB / "ni_score_test.csv", index=False)

    # Correzione congiunta di Holm-Bonferroni integrando i p-value di H1, H2 e H3
    print("\n" + "=" * 78)
    print(
        "CORREZIONE CONGIUNTA: H4 insieme a H1 (esaurimento), H2 (mieloidi), H3 (CAF)"
    )
    print("=" * 78)
    all_p = {"H4_ni_score": p}
    if args.previous_tests.exists():
        prev = pd.read_csv(args.previous_tests, index_col=0)
        for h in prev.index:
            all_p[h] = float(prev.loc[h, "p_perm"])
        print(f"[ok] caricati H1-H3 da {args.previous_tests.name}")
    else:
        print(
            f"[!] {args.previous_tests} assente: correggo solo H4 da sola (non ideale, esegui prima 16)."
        )

    adj = holm_bonferroni(all_p)
    print(f"\np-value grezzi: {all_p}")
    print("\np-value corretti (Holm-Bonferroni, famiglia completa):")
    for k, v in adj.items():
        sig = (
            "non definito"
            if v != v  # NaN: covariata costante, esclusa dalla famiglia
            else ("SIGNIFICATIVO" if v < 0.05 else "non significativo")
        )
        marker = "  [H4: questo script]" if k == "H4_ni_score" else ""
        print(f"  {k:20s} p_adj = {v:.4f}   {sig}{marker}")

    # il nome della colonna dichiara la dimensione della famiglia, cioe' i p definiti
    # (holm_bonferroni esclude i p NaN): 4 con il comando del README
    m = int(sum(np.isfinite(v) for v in all_p.values()))
    col = f"p_holm_family_of_{m}"
    pd.DataFrame({"p_raw": pd.Series(all_p), col: pd.Series(adj)}).rename_axis(
        "hypothesis"
    ).to_csv(TAB / "ni_score_holm_joint.csv")
    print(
        f"[ok] ni_score_holm_joint.csv: famiglia di {m} test con p definito su "
        f"{len(all_p)} (colonna {col})"
    )

    print("\n" + "=" * 78)
    print(f"n={len(df)} pazienti: come per H1-H3, questo e' un test a bassa potenza.")
    print("Un p<0.05 qui sarebbe un indizio (variabile clinica misurata dagli autori,")
    print("non derivata dal trascrittoma) ma non proverebbe causalita', e andrebbe")
    print("verificato su una coorte esterna prima di essere riportato come risultato.")
    print("=" * 78)

    # Generazione del grafico NI score vs tasso di errore colorato per stadio clinico TNM
    fig, ax = plt.subplots(figsize=(7, 5.5))
    stage_order = ["IA", "IB", "IIA", "IIB", "III"]
    cmap = plt.cm.viridis
    colors = [
        cmap(stage_order.index(s) / (len(stage_order) - 1))
        if s in stage_order
        else "gray"
        for s in df["Clinical Stage"]
    ]
    sc_ = ax.scatter(  # noqa: F841
        df["NI score"], df.error_rate, c=colors, s=160, edgecolor="k", zorder=5
    )
    for _, r in df.iterrows():
        ax.annotate(
            r.patient,
            (r["NI score"], r.error_rate),
            fontsize=8,
            xytext=(5, 5),
            textcoords="offset points",
        )
    ax.set_xlabel("NI score (invasione neurale, dallo studio originale)")
    ax.set_ylabel("tasso d'errore di annotazione (Tumor)")
    ax.set_title(f"NI score vs errore: rho={rho:+.3f}, p={p:.4f} (n={len(df)})")
    handles = [
        plt.Line2D(
            [0],
            [0],
            marker="o",
            color="w",
            markerfacecolor=cmap(i / (len(stage_order) - 1)),
            markersize=9,
            markeredgecolor="k",
            label=s,
        )
        for i, s in enumerate(stage_order)
    ]
    ax.legend(handles=handles, title="stadio clinico", fontsize=8, loc="best")
    fig.tight_layout()
    fig.savefig(FIG / "ni_score_vs_error.png", dpi=140)
    plt.close(fig)
    print(f"\n[ok] figura: results/figures/ni_score_vs_error.png")  # noqa: F541


if __name__ == "__main__":
    main()
